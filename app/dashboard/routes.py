from calendar import monthrange
from datetime import datetime, timedelta

from flask import (
    Blueprint,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    url_for
)

from flask_login import (
    current_user,
    login_required
)

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.member import Member
from app.models.group import Group
from app.models.program import Program
from app.models.reminder import Reminder
from app.models.campaign import Campaign
from app.services.campaign_service import CampaignService
from app.models.message import Message
from app.models.church import Church

dashboard_bp = Blueprint(
    "dashboard",
    __name__
)


# =========================================================
# DASHBOARD HOME
# =========================================================

@dashboard_bp.route("/")
@login_required
def home():

    church_id = current_user.church_id

    # =====================================================
    # MEMBERS
    # =====================================================

    member_count = db.session.scalar(
        db.select(
            db.func.count(Member.id)
        ).where(
            Member.church_id == church_id
        )
    ) or 0

    active_member_count = db.session.scalar(
        db.select(
            db.func.count(Member.id)
        ).where(
            Member.church_id == church_id,
            Member.is_active.is_(True)
        )
    ) or 0

    sms_opt_in_count = db.session.scalar(
        db.select(
            db.func.count(Member.id)
        ).where(
            Member.church_id == church_id,
            Member.is_active.is_(True),
            Member.sms_opt_in.is_(True)
        )
    ) or 0

    # =====================================================
    # PROGRAMS
    # =====================================================

    program_count = db.session.scalar(
        db.select(
            db.func.count(Program.id)
        ).where(
            Program.church_id == church_id
        )
    ) or 0

    active_program_count = db.session.scalar(
        db.select(
            db.func.count(Program.id)
        ).where(
            Program.church_id == church_id,
            Program.is_active.is_(True)
        )
    ) or 0

    # =====================================================
    # REMINDERS
    # =====================================================

    reminder_count = db.session.scalar(
        db.select(
            db.func.count(Reminder.id)
        ).where(
            Reminder.church_id == church_id
        )
    ) or 0

    active_reminder_count = db.session.scalar(
        db.select(
            db.func.count(Reminder.id)
        ).where(
            Reminder.church_id == church_id,
            Reminder.is_active.is_(True)
        )
    ) or 0

    scheduled_reminder_count = db.session.scalar(
        db.select(
            db.func.count(Reminder.id)
        ).where(
            Reminder.church_id == church_id,
            Reminder.status == "scheduled",
            Reminder.is_active.is_(True)
        )
    ) or 0

    # =====================================================
    # CAMPAIGNS
    # =====================================================

    campaign_count = db.session.scalar(
        db.select(
            db.func.count(Campaign.id)
        ).where(
            Campaign.church_id == church_id
        )
    ) or 0

    sent_campaign_count = db.session.scalar(
        db.select(
            db.func.count(Campaign.id)
        ).where(
            Campaign.church_id == church_id,
            Campaign.status == "sent"
        )
    ) or 0

    scheduled_campaign_count = db.session.scalar(
        db.select(
            db.func.count(Campaign.id)
        ).where(
            Campaign.church_id == church_id,
            Campaign.status == "scheduled"
        )
    ) or 0

    # =====================================================
    # MESSAGE HISTORY
    # =====================================================

    total_message_count = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == church_id
        )
    ) or 0

    sent_message_count = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == church_id,
            Message.status == "sent"
        )
    ) or 0

    failed_message_count = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == church_id,
            Message.status == "failed"
        )
    ) or 0

    pending_message_count = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == church_id,
            Message.status == "pending"
        )
    ) or 0

    # =====================================================
    # DELIVERY RATE
    # =====================================================

    delivery_rate = 0

    if total_message_count > 0:
        delivery_rate = round(
            (sent_message_count / total_message_count) * 100,
            1
        )

    # =====================================================
    # RECENT MESSAGE ACTIVITY
    # =====================================================

    recent_messages = db.session.scalars(
        db.select(Message)
        .where(
            Message.church_id == church_id
        )
        .order_by(
            Message.created_at.desc()
        )
        .limit(5)
    ).all()

    # =====================================================
    # UPCOMING REMINDERS
    # =====================================================

    upcoming_reminders = db.session.scalars(
        db.select(Reminder)
        .where(
            Reminder.church_id == church_id,
            Reminder.is_active.is_(True),
            Reminder.status == "scheduled",
            Reminder.next_run_at.is_not(None)
        )
        .order_by(
            Reminder.next_run_at.asc()
        )
        .limit(5)
    ).all()

    # =====================================================
    # UPCOMING CAMPAIGNS
    # =====================================================

    upcoming_campaigns = db.session.scalars(
        db.select(Campaign)
        .where(
            Campaign.church_id == church_id,
            Campaign.is_active.is_(True),
            Campaign.status == "scheduled",
            Campaign.scheduled_at.is_not(None)
        )
        .order_by(
            Campaign.scheduled_at.asc()
        )
        .limit(5)
    ).all()

    # =====================================================
    # RECENT PROGRAMS
    # =====================================================

    recent_programs = db.session.scalars(
        db.select(Program)
        .where(
            Program.church_id == church_id
        )
        .order_by(
            Program.id.desc()
        )
        .limit(5)
    ).all()

    # =====================================================
    # DASHBOARD
    # =====================================================

    return render_template(
        "dashboard/index.html",

        # Members
        member_count=member_count,
        active_member_count=active_member_count,
        sms_opt_in_count=sms_opt_in_count,

        # Programs
        program_count=program_count,
        active_program_count=active_program_count,

        # Reminders
        reminder_count=reminder_count,
        active_reminder_count=active_reminder_count,
        scheduled_reminder_count=scheduled_reminder_count,

        # Campaigns
        campaign_count=campaign_count,
        sent_campaign_count=sent_campaign_count,
        scheduled_campaign_count=scheduled_campaign_count,

        # Messages
        total_message_count=total_message_count,
        sent_message_count=sent_message_count,
        failed_message_count=failed_message_count,
        pending_message_count=pending_message_count,
        delivery_rate=delivery_rate,

        # Activity
        recent_messages=recent_messages,
        upcoming_reminders=upcoming_reminders,
        upcoming_campaigns=upcoming_campaigns,
        recent_programs=recent_programs
    )


# =========================================================
# MEMBERS
# =========================================================

@dashboard_bp.route("/members")
@login_required
def members():

    search = request.args.get(
        "search",
        ""
    ).strip()

    group_id = request.args.get(
        "group_id",
        type=int
    )

    query = db.select(Member).where(
        Member.church_id == current_user.church_id
    )

    if search:

        search_pattern = f"%{search}%"

        query = query.where(
            db.or_(
                Member.name.ilike(search_pattern),
                Member.phone.ilike(search_pattern),
                Member.email.ilike(search_pattern)
            )
        )

    if group_id:

        query = query.join(
            Member.groups
        ).where(
            Group.id == group_id
        )

    query = query.order_by(
        Member.is_active.desc(),
        Member.name.asc()
    )

    members_list = db.session.scalars(
        query
    ).unique().all()

    groups = db.session.scalars(
        db.select(Group)
        .where(
            Group.church_id == current_user.church_id,
            Group.is_active.is_(True)
        )
        .order_by(Group.name.asc())
    ).all()

    total_members = db.session.scalar(
        db.select(db.func.count(Member.id)).where(
            Member.church_id == current_user.church_id
        )
    ) or 0

    active_members = db.session.scalar(
        db.select(db.func.count(Member.id)).where(
            Member.church_id == current_user.church_id,
            Member.is_active.is_(True)
        )
    ) or 0

    inactive_members = db.session.scalar(
        db.select(db.func.count(Member.id)).where(
            Member.church_id == current_user.church_id,
            Member.is_active.is_(False)
        )
    ) or 0

    sms_members = db.session.scalar(
        db.select(db.func.count(Member.id)).where(
            Member.church_id == current_user.church_id,
            Member.sms_opt_in.is_(True),
            Member.is_active.is_(True)
        )
    ) or 0

    return render_template(
        "members/index.html",
        members=members_list,
        groups=groups,
        search=search,
        selected_group_id=group_id,
        total_members=total_members,
        active_members=active_members,
        inactive_members=inactive_members,
        sms_members=sms_members
    )


# =========================================================
# MEMBER DETAILS
# =========================================================

@dashboard_bp.route(
    "/members/<int:member_id>"
)
@login_required
def member_details(member_id):

    member = db.session.scalar(
        db.select(Member).where(
            Member.id == member_id,
            Member.church_id == current_user.church_id
        )
    )

    if member is None:

        flash(
            "Member not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.members")
        )

    return render_template(
        "members/details.html",
        member=member
    )


# =========================================================
# EDIT MEMBER
# =========================================================

@dashboard_bp.route(
    "/members/<int:member_id>/edit",
    methods=["GET", "POST"]
)
@login_required
def edit_member(member_id):

    member = db.session.scalar(
        db.select(Member).where(
            Member.id == member_id,
            Member.church_id == current_user.church_id
        )
    )

    if member is None:

        flash(
            "Member not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.members")
        )

    groups = db.session.scalars(
        db.select(Group)
        .where(
            Group.church_id == current_user.church_id,
            Group.is_active.is_(True)
        )
        .order_by(Group.name.asc())
    ).all()

    if request.method == "POST":

        name = request.form.get(
            "name",
            ""
        ).strip()

        phone = request.form.get(
            "phone",
            ""
        ).strip()

        email = request.form.get(
            "email",
            ""
        ).strip() or None

        group_ids = request.form.getlist(
            "groups"
        )

        sms_opt_in = (
            "sms_opt_in"
            in request.form
        )

        whatsapp_opt_in = (
            "whatsapp_opt_in"
            in request.form
        )

        if not name:

            flash(
                "Member name is required.",
                "error"
            )

            return render_template(
                "members/edit.html",
                member=member,
                groups=groups
            )

        if not phone:

            flash(
                "Phone number is required.",
                "error"
            )

            return render_template(
                "members/edit.html",
                member=member,
                groups=groups
            )

        member.name = name
        member.phone = phone
        member.email = email
        member.sms_opt_in = sms_opt_in
        member.whatsapp_opt_in = whatsapp_opt_in

        selected_groups = []

        for group_id in group_ids:

            group = db.session.scalar(
                db.select(Group).where(
                    Group.id == int(group_id),
                    Group.church_id == current_user.church_id
                )
            )

            if group:
                selected_groups.append(group)

        member.groups = selected_groups

        db.session.commit()

        flash(
            f"{member.name} was updated successfully.",
            "success"
        )

        return redirect(
            url_for(
                "dashboard.member_details",
                member_id=member.id
            )
        )

    return render_template(
        "members/edit.html",
        member=member,
        groups=groups
    )


# =========================================================
# ADD MEMBER
# =========================================================

@dashboard_bp.route(
    "/members/add",
    methods=["GET", "POST"]
)
@login_required
def add_member():

    groups = db.session.scalars(
        db.select(Group)
        .where(
            Group.church_id == current_user.church_id,
            Group.is_active.is_(True)
        )
        .order_by(Group.name.asc())
    ).all()

    if request.method == "POST":

        name = request.form.get(
            "name",
            ""
        ).strip()

        phone = request.form.get(
            "phone",
            ""
        ).strip()

        email = request.form.get(
            "email",
            ""
        ).strip() or None

        group_ids = request.form.getlist(
            "groups"
        )

        sms_opt_in = (
            "sms_opt_in"
            in request.form
        )

        whatsapp_opt_in = (
            "whatsapp_opt_in"
            in request.form
        )

        if not name:

            flash(
                "Member name is required.",
                "error"
            )

            return render_template(
                "members/add.html",
                groups=groups
            )

        if not phone:

            flash(
                "Phone number is required.",
                "error"
            )

            return render_template(
                "members/add.html",
                groups=groups
            )

        member = Member(
            church_id=current_user.church_id,
            name=name,
            phone=phone,
            email=email,
            sms_opt_in=sms_opt_in,
            whatsapp_opt_in=whatsapp_opt_in,
            is_active=True
        )

        selected_groups = []

        for group_id in group_ids:

            group = db.session.scalar(
                db.select(Group).where(
                    Group.id == int(group_id),
                    Group.church_id == current_user.church_id
                )
            )

            if group:
                selected_groups.append(group)

        member.groups = selected_groups

        db.session.add(member)

        db.session.commit()

        flash(
            f"{member.name} was added successfully.",
            "success"
        )

        return redirect(
            url_for(
                "dashboard.member_details",
                member_id=member.id
            )
        )

    return render_template(
        "members/add.html",
        groups=groups
    )


# =========================================================
# TOGGLE MEMBER STATUS
# =========================================================

@dashboard_bp.route(
    "/members/<int:member_id>/toggle-status",
    methods=["POST"]
)
@login_required
def toggle_member_status(member_id):

    member = db.session.scalar(
        db.select(Member).where(
            Member.id == member_id,
            Member.church_id == current_user.church_id
        )
    )

    if member is None:

        flash(
            "Member not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.members")
        )

    member.is_active = not member.is_active

    db.session.commit()

    if member.is_active:

        flash(
            f"{member.name} has been activated.",
            "success"
        )

    else:

        flash(
            f"{member.name} has been deactivated.",
            "success"
        )

    return redirect(
        url_for(
            "dashboard.member_details",
            member_id=member.id
        )
    )


# =========================================================
# DELETE MEMBER
# =========================================================

@dashboard_bp.route(
    "/members/<int:member_id>/delete",
    methods=["POST"]
)
@login_required
def delete_member(member_id):

    member = db.session.scalar(
        db.select(Member).where(
            Member.id == member_id,
            Member.church_id == current_user.church_id
        )
    )

    if member is None:

        flash(
            "Member not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.members")
        )

    member_name = member.name

    db.session.delete(member)

    db.session.commit()

    flash(
        f"{member_name} was deleted successfully.",
        "success"
    )

    return redirect(
        url_for("dashboard.members")
    )


# =========================================================
# IMPORT MEMBERS
# =========================================================

@dashboard_bp.route(
    "/members/import",
    methods=["GET", "POST"]
)
@login_required
def import_members():

    if request.method == "POST":

        file = request.files.get(
            "file"
        )

        if not file or not file.filename:

            flash(
                "Please select an Excel file.",
                "error"
            )

            return redirect(
                url_for("dashboard.import_members")
            )

        try:

            workbook = load_workbook(
                file,
                data_only=True
            )

            worksheet = workbook.active

            rows = list(
                worksheet.iter_rows(
                    values_only=True
                )
            )

            if not rows:

                flash(
                    "The Excel file is empty.",
                    "error"
                )

                return redirect(
                    url_for("dashboard.import_members")
                )

            headers = [
                str(cell).strip().lower()
                if cell is not None
                else ""
                for cell in rows[0]
            ]

            required_headers = [
                "name",
                "phone"
            ]

            for header in required_headers:

                if header not in headers:

                    flash(
                        f"The Excel file must contain a '{header}' column.",
                        "error"
                    )

                    return redirect(
                        url_for("dashboard.import_members")
                    )

            name_index = headers.index(
                "name"
            )

            phone_index = headers.index(
                "phone"
            )

            email_index = (
                headers.index("email")
                if "email" in headers
                else None
            )

            group_index = (
                headers.index("group")
                if "group" in headers
                else None
            )

            imported_count = 0
            skipped_count = 0

            for row in rows[1:]:

                if not row:
                    continue

                name = (
                    str(row[name_index]).strip()
                    if len(row) > name_index
                    and row[name_index] is not None
                    else ""
                )

                phone = (
                    str(row[phone_index]).strip()
                    if len(row) > phone_index
                    and row[phone_index] is not None
                    else ""
                )

                if not name or not phone:

                    skipped_count += 1
                    continue

                email = None

                if (
                    email_index is not None
                    and len(row) > email_index
                    and row[email_index] is not None
                ):

                    email = str(
                        row[email_index]
                    ).strip() or None

                member = Member(
                    church_id=current_user.church_id,
                    name=name,
                    phone=phone,
                    email=email,
                    is_active=True,
                    sms_opt_in=True,
                    whatsapp_opt_in=False
                )

                if group_index is not None:

                    if (
                        len(row) > group_index
                        and row[group_index] is not None
                    ):

                        group_name = str(
                            row[group_index]
                        ).strip()

                        if group_name:

                            group = db.session.scalar(
                                db.select(Group).where(
                                    Group.church_id == current_user.church_id,
                                    db.func.lower(Group.name)
                                    == group_name.lower()
                                )
                            )

                            if group:

                                member.groups.append(
                                    group
                                )

                db.session.add(
                    member
                )

                imported_count += 1

            db.session.commit()

            flash(
                f"{imported_count} members imported successfully. "
                f"{skipped_count} rows skipped.",
                "success"
            )

            return redirect(
                url_for("dashboard.members")
            )

        except Exception as error:

            db.session.rollback()

            flash(
                f"Could not import the Excel file: {error}",
                "error"
            )

            return redirect(
                url_for("dashboard.import_members")
            )

    return render_template(
        "members/import.html"
    )


# =========================================================
# DOWNLOAD MEMBER IMPORT TEMPLATE
# =========================================================

@dashboard_bp.route(
    "/members/import/template"
)
@login_required
def download_member_template():

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Members"

    worksheet.append([
        "name",
        "phone",
        "email",
        "group"
    ])

    worksheet.append([
        "Example Member",
        "0590000000",
        "example@email.com",
        "Youth"
    ])

    file_path = "member_import_template.xlsx"

    workbook.save(
        file_path
    )

    return send_file(
        file_path,
        as_attachment=True,
        download_name="member_import_template.xlsx"
    )


# =========================================================
# PROGRAMS / SERVICES
# =========================================================

@dashboard_bp.route("/programs")
@login_required
def programs():

    programs = db.session.scalars(
        db.select(Program)
        .where(
            Program.church_id == current_user.church_id
        )
        .order_by(
            Program.is_active.desc(),
            Program.name.asc()
        )
    ).all()

    active_programs = sum(
        1
        for program in programs
        if program.is_active
    )

    inactive_programs = sum(
        1
        for program in programs
        if not program.is_active
    )

    return render_template(
        "programs/index.html",
        programs=programs,
        active_programs=active_programs,
        inactive_programs=inactive_programs
    )


# =========================================================
# PROGRAM DETAILS
# =========================================================

@dashboard_bp.route(
    "/programs/<int:program_id>"
)
@login_required
def program_details(program_id):

    program = db.session.scalar(
        db.select(Program).where(
            Program.id == program_id,
            Program.church_id == current_user.church_id
        )
    )

    if program is None:

        flash(
            "Program not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.programs")
        )

    return render_template(
        "programs/details.html",
        program=program
    )


# =========================================================
# ADD PROGRAM
# =========================================================

@dashboard_bp.route(
    "/programs/add",
    methods=["GET", "POST"]
)
@login_required
def add_program():

    if request.method == "POST":

        name = request.form.get(
            "name",
            ""
        ).strip()

        description = request.form.get(
            "description",
            ""
        ).strip() or None

        day_of_week = request.form.get(
            "day_of_week",
            ""
        ).strip() or None

        start_time_value = request.form.get(
            "start_time",
            ""
        ).strip()

        location = request.form.get(
            "location",
            ""
        ).strip() or None

        if not name:

            flash(
                "Program name is required.",
                "error"
            )

            return render_template(
                "programs/add.html"
            )

        start_time = None

        if start_time_value:

            try:

                start_time = datetime.strptime(
                    start_time_value,
                    "%H:%M"
                ).time()

            except ValueError:

                flash(
                    "Please enter a valid start time.",
                    "error"
                )

                return render_template(
                    "programs/add.html"
                )

        program = Program(
            church_id=current_user.church_id,
            name=name,
            description=description,
            day_of_week=day_of_week,
            start_time=start_time,
            location=location,
            is_active=True
        )

        db.session.add(
            program
        )

        db.session.commit()

        flash(
            f"{name} was added successfully.",
            "success"
        )

        return redirect(
            url_for("dashboard.programs")
        )

    return render_template(
        "programs/add.html"
    )


# =========================================================
# EDIT PROGRAM
# =========================================================

@dashboard_bp.route(
    "/programs/<int:program_id>/edit",
    methods=["GET", "POST"]
)
@login_required
def edit_program(program_id):

    program = db.session.scalar(
        db.select(Program).where(
            Program.id == program_id,
            Program.church_id == current_user.church_id
        )
    )

    if program is None:

        flash(
            "Program not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.programs")
        )

    if request.method == "POST":

        name = request.form.get(
            "name",
            ""
        ).strip()

        description = request.form.get(
            "description",
            ""
        ).strip() or None

        day_of_week = request.form.get(
            "day_of_week",
            ""
        ).strip() or None

        start_time_value = request.form.get(
            "start_time",
            ""
        ).strip()

        location = request.form.get(
            "location",
            ""
        ).strip() or None

        is_active = (
            "is_active"
            in request.form
        )

        if not name:

            flash(
                "Program name is required.",
                "error"
            )

            return render_template(
                "programs/edit.html",
                program=program
            )

        start_time = None

        if start_time_value:

            try:

                start_time = datetime.strptime(
                    start_time_value,
                    "%H:%M"
                ).time()

            except ValueError:

                flash(
                    "Please enter a valid start time.",
                    "error"
                )

                return render_template(
                    "programs/edit.html",
                    program=program
                )

        program.name = name
        program.description = description
        program.day_of_week = day_of_week
        program.start_time = start_time
        program.location = location
        program.is_active = is_active

        db.session.commit()

        flash(
            f"{program.name} was updated successfully.",
            "success"
        )

        return redirect(
            url_for(
                "dashboard.program_details",
                program_id=program.id
            )
        )

    return render_template(
        "programs/edit.html",
        program=program
    )


# =========================================================
# TOGGLE PROGRAM STATUS
# =========================================================

@dashboard_bp.route(
    "/programs/<int:program_id>/toggle-status",
    methods=["POST"]
)
@login_required
def toggle_program_status(program_id):

    program = db.session.scalar(
        db.select(Program).where(
            Program.id == program_id,
            Program.church_id == current_user.church_id
        )
    )

    if program is None:

        flash(
            "Program not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.programs")
        )

    program.is_active = not program.is_active

    db.session.commit()

    if program.is_active:

        flash(
            f"{program.name} has been activated.",
            "success"
        )

    else:

        flash(
            f"{program.name} has been deactivated.",
            "success"
        )

    return redirect(
        url_for(
            "dashboard.program_details",
            program_id=program.id
        )
    )


# =========================================================
# DELETE PROGRAM
# =========================================================

@dashboard_bp.route(
    "/programs/<int:program_id>/delete",
    methods=["POST"]
)
@login_required
def delete_program(program_id):

    program = db.session.scalar(
        db.select(Program).where(
            Program.id == program_id,
            Program.church_id == current_user.church_id
        )
    )

    if program is None:

        flash(
            "Program not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.programs")
        )

    program_name = program.name

    db.session.delete(
        program
    )

    db.session.commit()

    flash(
        f"{program_name} was deleted successfully.",
        "success"
    )

    return redirect(
        url_for("dashboard.programs")
    )


# =========================================================
# DATABASE TEST
# =========================================================

@dashboard_bp.route(
    "/database-test"
)
@login_required
def database_test():

    try:

        result = db.session.execute(
            db.text("SELECT 1")
        )

        result.scalar()

        return "Database connection successful."

    except Exception as error:

        return (
            "Database connection failed: "
            f"{error}"
        )


# =========================================================
# REMINDER HELPERS
# =========================================================

def calculate_next_run(
    scheduled_at,
    recurrence_type,
    recurrence_interval=1
):
    """
    Calculate the next occurrence of a reminder.

    Supported recurrence types:

        one_time
        daily
        weekly
        monthly
        yearly
        interval_days
        interval_weeks
        interval_months
        interval_years
    """

    if not scheduled_at:
        return None

    interval = recurrence_interval or 1

    if interval < 1:
        interval = 1

    if recurrence_type == "one_time":
        return None

    if recurrence_type == "daily":

        return scheduled_at + timedelta(
            days=1
        )

    if recurrence_type == "weekly":

        return scheduled_at + timedelta(
            weeks=1
        )

    if recurrence_type == "monthly":

        total_months = (
            scheduled_at.year * 12
            + scheduled_at.month - 1
            + 1
        )

        year = total_months // 12
        month = total_months % 12 + 1

        last_day = monthrange(
            year,
            month
        )[1]

        day = min(
            scheduled_at.day,
            last_day
        )

        return scheduled_at.replace(
            year=year,
            month=month,
            day=day
        )

    if recurrence_type == "yearly":

        try:

            return scheduled_at.replace(
                year=scheduled_at.year + 1
            )

        except ValueError:

            return scheduled_at.replace(
                year=scheduled_at.year + 1,
                day=28
            )

    if recurrence_type == "interval_days":

        return scheduled_at + timedelta(
            days=interval
        )

    if recurrence_type == "interval_weeks":

        return scheduled_at + timedelta(
            weeks=interval
        )

    if recurrence_type == "interval_months":

        total_months = (
            scheduled_at.year * 12
            + scheduled_at.month - 1
            + interval
        )

        year = total_months // 12
        month = total_months % 12 + 1

        last_day = monthrange(
            year,
            month
        )[1]

        day = min(
            scheduled_at.day,
            last_day
        )

        return scheduled_at.replace(
            year=year,
            month=month,
            day=day
        )

    if recurrence_type == "interval_years":

        target_year = (
            scheduled_at.year
            + interval
        )

        try:

            return scheduled_at.replace(
                year=target_year
            )

        except ValueError:

            return scheduled_at.replace(
                year=target_year,
                day=28
            )

    return None


def calculate_next_future_run(
    scheduled_at,
    recurrence_type,
    recurrence_interval,
    now=None
):
    """
    Calculate the next recurring run that occurs after now.

    This is useful when an existing recurring reminder is edited
    and its original scheduled time is already in the past.
    """

    if not scheduled_at:
        return None

    if recurrence_type == "one_time":
        return None

    if now is None:
        now = datetime.utcnow()

    next_run = calculate_next_run(
        scheduled_at,
        recurrence_type,
        recurrence_interval
    )

    safety_counter = 0

    while (
        next_run
        and next_run <= now
        and safety_counter < 10000
    ):

        next_run = calculate_next_run(
            next_run,
            recurrence_type,
            recurrence_interval
        )

        safety_counter += 1

    return next_run


def get_reminder_form_data():
    """
    Load programs and groups belonging only to the
    currently logged-in user's church.
    """

    programs = db.session.scalars(
        db.select(Program)
        .where(
            Program.church_id == current_user.church_id,
            Program.is_active.is_(True)
        )
        .order_by(
            Program.name.asc()
        )
    ).all()

    groups = db.session.scalars(
        db.select(Group)
        .where(
            Group.church_id == current_user.church_id,
            Group.is_active.is_(True)
        )
        .order_by(
            Group.name.asc()
        )
    ).all()

    return programs, groups


# =========================================================
# ADD REMINDER
# =========================================================

@dashboard_bp.route(
    "/reminders/add",
    methods=["GET", "POST"]
)
@login_required
def add_reminder():

    programs, groups = get_reminder_form_data()

    if request.method == "POST":

        title = request.form.get(
            "title",
            ""
        ).strip()

        message = request.form.get(
            "message",
            ""
        ).strip()

        program_id = request.form.get(
            "program_id",
            ""
        ).strip()

        audience_type = request.form.get(
            "audience_type",
            "all"
        ).strip()

        scheduled_at_raw = request.form.get(
            "scheduled_at",
            ""
        ).strip()

        recurrence_type = request.form.get(
            "recurrence_type",
            "one_time"
        ).strip()

        recurrence_interval_raw = request.form.get(
            "recurrence_interval",
            "1"
        ).strip()

        recurrence_end_date_raw = request.form.get(
            "recurrence_end_date",
            ""
        ).strip()

        group_ids = request.form.getlist(
            "group_ids"
        )

        if not title:

            flash(
                "Please enter a reminder title.",
                "error"
            )

            return render_template(
                "reminders/add.html",
                programs=programs,
                groups=groups
            )

        if not message:

            flash(
                "Please enter the SMS message.",
                "error"
            )

            return render_template(
                "reminders/add.html",
                programs=programs,
                groups=groups
            )

        if not scheduled_at_raw:

            flash(
                "Please select a date and time.",
                "error"
            )

            return render_template(
                "reminders/add.html",
                programs=programs,
                groups=groups
            )

        try:

            scheduled_at = datetime.strptime(
                scheduled_at_raw,
                "%Y-%m-%dT%H:%M"
            )

        except ValueError:

            flash(
                "Invalid date and time format.",
                "error"
            )

            return render_template(
                "reminders/add.html",
                programs=programs,
                groups=groups
            )

        allowed_recurrence_types = {
            "one_time",
            "daily",
            "weekly",
            "monthly",
            "yearly",
            "interval_days",
            "interval_weeks",
            "interval_months",
            "interval_years"
        }

        if recurrence_type not in allowed_recurrence_types:

            flash(
                "Invalid recurrence option selected.",
                "error"
            )

            return render_template(
                "reminders/add.html",
                programs=programs,
                groups=groups
            )

        try:

            recurrence_interval = int(
                recurrence_interval_raw or 1
            )

        except ValueError:

            flash(
                "Recurrence interval must be a number.",
                "error"
            )

            return render_template(
                "reminders/add.html",
                programs=programs,
                groups=groups
            )

        if recurrence_interval < 1:

            flash(
                "Recurrence interval must be at least 1.",
                "error"
            )

            return render_template(
                "reminders/add.html",
                programs=programs,
                groups=groups
            )

        recurrence_end_date = None

        if recurrence_end_date_raw:

            try:

                recurrence_end_date = datetime.strptime(
                    recurrence_end_date_raw,
                    "%Y-%m-%d"
                )

            except ValueError:

                flash(
                    "Invalid recurrence end date.",
                    "error"
                )

                return render_template(
                    "reminders/add.html",
                    programs=programs,
                    groups=groups
                )

            if recurrence_end_date < scheduled_at:

                flash(
                    "The recurrence end date cannot be before the scheduled date.",
                    "error"
                )

                return render_template(
                    "reminders/add.html",
                    programs=programs,
                    groups=groups
                )

        if recurrence_type == "one_time":

            recurrence_interval = 1
            recurrence_end_date = None

        allowed_audiences = {
            "all",
            "groups"
        }

        if audience_type not in allowed_audiences:

            flash(
                "Invalid audience selected.",
                "error"
            )

            return render_template(
                "reminders/add.html",
                programs=programs,
                groups=groups
            )

        selected_groups = []

        if audience_type == "groups":

            if not group_ids:

                flash(
                    "Please select at least one group.",
                    "error"
                )

                return render_template(
                    "reminders/add.html",
                    programs=programs,
                    groups=groups
                )

            selected_group_ids = []

            for group_id in group_ids:

                try:

                    selected_group_ids.append(
                        int(group_id)
                    )

                except ValueError:

                    continue

            if not selected_group_ids:

                flash(
                    "Invalid group selection.",
                    "error"
                )

                return render_template(
                    "reminders/add.html",
                    programs=programs,
                    groups=groups
                )

            selected_groups = db.session.scalars(
                db.select(Group)
                .where(
                    Group.id.in_(
                        selected_group_ids
                    ),
                    Group.church_id == current_user.church_id,
                    Group.is_active.is_(True)
                )
            ).all()

            if len(selected_groups) != len(
                set(selected_group_ids)
            ):

                flash(
                    "One or more selected groups are invalid.",
                    "error"
                )

                return render_template(
                    "reminders/add.html",
                    programs=programs,
                    groups=groups
                )

        program = None

        if program_id:

            try:

                program_id_int = int(
                    program_id
                )

            except ValueError:

                flash(
                    "Invalid program selected.",
                    "error"
                )

                return render_template(
                    "reminders/add.html",
                    programs=programs,
                    groups=groups
                )

            program = db.session.scalar(
                db.select(Program)
                .where(
                    Program.id == program_id_int,
                    Program.church_id == current_user.church_id,
                    Program.is_active.is_(True)
                )
            )

            if not program:

                flash(
                    "The selected program is invalid.",
                    "error"
                )

                return render_template(
                    "reminders/add.html",
                    programs=programs,
                    groups=groups
                )

        next_run_at = calculate_next_run(
            scheduled_at=scheduled_at,
            recurrence_type=recurrence_type,
            recurrence_interval=recurrence_interval
        )

        if (
            next_run_at
            and recurrence_end_date
            and next_run_at > recurrence_end_date
        ):

            next_run_at = None

        reminder = Reminder(
            church_id=current_user.church_id,

            program_id=(
                program.id
                if program
                else None
            ),

            title=title,

            message=message,

            audience_type=audience_type,

            channel="sms",

            scheduled_at=scheduled_at,

            recurrence_type=recurrence_type,

            recurrence_interval=recurrence_interval,

            recurrence_end_date=recurrence_end_date,

            next_run_at=next_run_at,

            status="scheduled",

            is_active=True,

            total_sent=0
        )

        if audience_type == "groups":

            reminder.groups = selected_groups

        db.session.add(
            reminder
        )

        db.session.commit()

        flash(
            "Reminder created successfully.",
            "success"
        )

        return redirect(
            url_for(
                "dashboard.reminders"
            )
        )

    return render_template(
        "reminders/add.html",
        programs=programs,
        groups=groups
    )


# =========================================================
# CAMPAIGNS
# =========================================================

@dashboard_bp.route("/campaigns")
@login_required
def campaigns():
    campaigns = db.session.scalars(
        db.select(Campaign)
        .where(Campaign.church_id == current_user.church_id)
        .order_by(Campaign.created_at.desc())
    ).all()

    return render_template(
        "campaigns/index.html",
        campaigns=campaigns
    )


@dashboard_bp.route("/campaigns/new", methods=["GET", "POST"])
@login_required
def add_campaign():

    groups = db.session.scalars(
        db.select(Group)
        .where(
            Group.church_id == current_user.church_id,
            Group.is_active.is_(True)
        )
        .order_by(Group.name.asc())
    ).all()

    if request.method == "POST":

        title = request.form.get("title", "").strip()
        message = request.form.get("message", "").strip()
        audience_type = request.form.get("audience_type", "all")
        channel = request.form.get("channel", "sms")
        scheduled_at_raw = request.form.get("scheduled_at", "").strip()

        if not title:
            flash("Campaign title is required.", "error")

            return render_template(
                "campaigns/create.html",
                groups=groups
            )

        if not message:
            flash("Campaign message is required.", "error")

            return render_template(
                "campaigns/create.html",
                groups=groups
            )

        if channel != "sms":
            flash("SMS is currently the available campaign channel.", "error")

            return render_template(
                "campaigns/create.html",
                groups=groups
            )

        selected_group_ids = request.form.getlist("group_ids")

        selected_groups = []

        if audience_type == "selected_groups":

            if not selected_group_ids:
                flash("Please select at least one group.", "error")

                return render_template(
                    "campaigns/create.html",
                    groups=groups
                )

            selected_groups = db.session.scalars(
                db.select(Group)
                .where(
                    Group.church_id == current_user.church_id,
                    Group.id.in_(selected_group_ids),
                    Group.is_active.is_(True)
                )
            ).all()

            if not selected_groups:
                flash("The selected groups could not be found.", "error")

                return render_template(
                    "campaigns/create.html",
                    groups=groups
                )

        scheduled_at = None

        if scheduled_at_raw:

            try:
                scheduled_at = datetime.strptime(
                    scheduled_at_raw,
                    "%Y-%m-%dT%H:%M"
                )

            except ValueError:
                flash("Please enter a valid scheduled date and time.", "error")

                return render_template(
                    "campaigns/create.html",
                    groups=groups
                )

        # Find eligible recipients.
        if audience_type == "all":

            recipients = db.session.scalars(
                db.select(Member)
                .where(
                    Member.church_id == current_user.church_id,
                    Member.is_active.is_(True),
                    Member.sms_opt_in.is_(True)
                )
            ).all()

        else:

            recipient_ids = set()

            for group in selected_groups:
                for member in group.members:

                    if (
                        member.church_id == current_user.church_id
                        and member.is_active
                        and member.sms_opt_in
                    ):
                        recipient_ids.add(member.id)

            if recipient_ids:

                recipients = db.session.scalars(
                    db.select(Member)
                    .where(
                        Member.church_id == current_user.church_id,
                        Member.id.in_(recipient_ids),
                        Member.is_active.is_(True),
                        Member.sms_opt_in.is_(True)
                    )
                ).all()

            else:
                recipients = []

        total_recipients = len(recipients)

        campaign = Campaign(
            church_id=current_user.church_id,
            title=title,
            message=message,
            audience_type=audience_type,
            channel=channel,
            scheduled_at=scheduled_at,
            status="scheduled" if scheduled_at else "draft",
            is_active=True,
            total_recipients=total_recipients,
            total_sent=0,
            total_failed=0
        )

        if selected_groups:
            campaign.groups = selected_groups

        db.session.add(campaign)
        db.session.commit()

        if scheduled_at:
            flash(
                f"Campaign scheduled successfully for {scheduled_at.strftime('%d %b %Y, %I:%M %p')}.",
                "success"
            )
        else:
            flash(
                "Campaign saved as a draft successfully.",
                "success"
            )

        return redirect(
            url_for("dashboard.campaigns")
        )

    return render_template(
        "campaigns/create.html",
        groups=groups
    )

# =========================================================
# CAMPAIGN DETAILS
# =========================================================

@dashboard_bp.route(
    "/campaigns/<int:campaign_id>"
)
@login_required
def campaign_details(campaign_id):

    campaign = db.session.scalar(
        db.select(Campaign).where(
            Campaign.id == campaign_id,
            Campaign.church_id == current_user.church_id
        )
    )

    if campaign is None:

        flash(
            "Campaign not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.campaigns")
        )

    return render_template(
        "campaigns/details.html",
        campaign=campaign
    )

# =========================================================
# EDIT CAMPAIGN
# =========================================================

@dashboard_bp.route(
    "/campaigns/<int:campaign_id>/edit",
    methods=["GET", "POST"]
)
@login_required
def edit_campaign(campaign_id):

    campaign = db.session.scalar(
        db.select(Campaign).where(
            Campaign.id == campaign_id,
            Campaign.church_id == current_user.church_id
        )
    )

    if campaign is None:

        flash(
            "Campaign not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.campaigns")
        )

    # Sent campaigns should remain as history.
    if campaign.status in ["sending", "sent"]:

        flash(
            "This campaign can no longer be edited.",
            "error"
        )

        return redirect(
            url_for(
                "dashboard.campaign_details",
                campaign_id=campaign.id
            )
        )

    groups = db.session.scalars(
        db.select(Group)
        .where(
            Group.church_id == current_user.church_id,
            Group.is_active.is_(True)
        )
        .order_by(Group.name.asc())
    ).all()

    if request.method == "POST":

        title = request.form.get(
            "title",
            ""
        ).strip()

        message = request.form.get(
            "message",
            ""
        ).strip()

        audience_type = request.form.get(
            "audience_type",
            "all"
        )

        channel = request.form.get(
            "channel",
            "sms"
        )

        scheduled_at_raw = request.form.get(
            "scheduled_at",
            ""
        ).strip()

        selected_group_ids = request.form.getlist(
            "group_ids"
        )

        # -------------------------------------------------
        # BASIC VALIDATION
        # -------------------------------------------------

        if not title:

            flash(
                "Campaign title is required.",
                "error"
            )

            return render_template(
                "campaigns/edit.html",
                campaign=campaign,
                groups=groups
            )

        if not message:

            flash(
                "Campaign message is required.",
                "error"
            )

            return render_template(
                "campaigns/edit.html",
                campaign=campaign,
                groups=groups
            )

        if channel != "sms":

            flash(
                "SMS is currently the available campaign channel.",
                "error"
            )

            return render_template(
                "campaigns/edit.html",
                campaign=campaign,
                groups=groups
            )

        # -------------------------------------------------
        # GROUP VALIDATION
        # -------------------------------------------------

        selected_groups = []

        if audience_type == "selected_groups":

            if not selected_group_ids:

                flash(
                    "Please select at least one group.",
                    "error"
                )

                return render_template(
                    "campaigns/edit.html",
                    campaign=campaign,
                    groups=groups
                )

            selected_groups = db.session.scalars(
                db.select(Group)
                .where(
                    Group.church_id == current_user.church_id,
                    Group.id.in_(selected_group_ids),
                    Group.is_active.is_(True)
                )
            ).all()

            if not selected_groups:

                flash(
                    "The selected groups could not be found.",
                    "error"
                )

                return render_template(
                    "campaigns/edit.html",
                    campaign=campaign,
                    groups=groups
                )

        # -------------------------------------------------
        # SCHEDULE
        # -------------------------------------------------

        scheduled_at = None

        if scheduled_at_raw:

            try:

                scheduled_at = datetime.strptime(
                    scheduled_at_raw,
                    "%Y-%m-%dT%H:%M"
                )

            except ValueError:

                flash(
                    "Please enter a valid scheduled date and time.",
                    "error"
                )

                return render_template(
                    "campaigns/edit.html",
                    campaign=campaign,
                    groups=groups
                )

            if scheduled_at <= datetime.utcnow():

                flash(
                    "Scheduled time must be in the future.",
                    "error"
                )

                return render_template(
                    "campaigns/edit.html",
                    campaign=campaign,
                    groups=groups
                )

        # -------------------------------------------------
        # UPDATE
        # -------------------------------------------------

        campaign.title = title
        campaign.message = message
        campaign.audience_type = audience_type
        campaign.channel = channel
        campaign.scheduled_at = scheduled_at

        campaign.total_sent = 0
        campaign.total_failed = 0
        campaign.sent_at = None
        campaign.failure_reason = None

        campaign.status = (
            "scheduled"
            if scheduled_at
            else "draft"
        )

        campaign.groups = selected_groups

        db.session.commit()

        flash(
            "Campaign updated successfully.",
            "success"
        )

        return redirect(
            url_for(
                "dashboard.campaign_details",
                campaign_id=campaign.id
            )
        )

    return render_template(
        "campaigns/edit.html",
        campaign=campaign,
        groups=groups
    )


# =========================================================
# SEND CAMPAIGN NOW
# =========================================================

@dashboard_bp.route(
    "/campaigns/<int:campaign_id>/send",
    methods=["POST"]
)
@login_required
def send_campaign(campaign_id):

    campaign = db.session.scalar(
        db.select(Campaign).where(
            Campaign.id == campaign_id,
            Campaign.church_id == current_user.church_id
        )
    )

    if campaign is None:

        flash(
            "Campaign not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.campaigns")
        )

    if campaign.status == "sending":

        flash(
            "This campaign is already being processed.",
            "error"
        )

        return redirect(
            url_for(
                "dashboard.campaign_details",
                campaign_id=campaign.id
            )
        )

    if campaign.status == "sent":

        flash(
            "This campaign has already been sent.",
            "error"
        )

        return redirect(
            url_for(
                "dashboard.campaign_details",
                campaign_id=campaign.id
            )
        )

    result = CampaignService.process_campaign(
        campaign
    )

    if result["status"] == "sent":

        if result["failed"] > 0:

            flash(
                f"Campaign completed: "
                f"{result['sent']} sent, "
                f"{result['failed']} failed.",
                "warning"
            )

        else:

            flash(
                f"Campaign sent successfully to "
                f"{result['sent']} recipient(s).",
                "success"
            )

    else:

        flash(
            result["failure_reason"]
            or "Campaign sending failed.",
            "error"
        )

    return redirect(
        url_for(
            "dashboard.campaign_details",
            campaign_id=campaign.id
        )
    )


# =========================================================
# DELETE CAMPAIGN
# =========================================================

@dashboard_bp.route(
    "/campaigns/<int:campaign_id>/delete",
    methods=["POST"]
)
@login_required
def delete_campaign(campaign_id):

    campaign = db.session.scalar(
        db.select(Campaign).where(
            Campaign.id == campaign_id,
            Campaign.church_id == current_user.church_id
        )
    )

    if campaign is None:

        flash(
            "Campaign not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.campaigns")
        )

    if campaign.status == "sending":

        flash(
            "A campaign currently being sent cannot be deleted.",
            "error"
        )

        return redirect(
            url_for(
                "dashboard.campaign_details",
                campaign_id=campaign.id
            )
        )

    db.session.delete(campaign)
    db.session.commit()

    flash(
        "Campaign deleted successfully.",
        "success"
    )

    return redirect(
        url_for("dashboard.campaigns")
    )

# REMINDERS LIST
# =========================================================

@dashboard_bp.route(
    "/reminders"
)
@login_required
def reminders():

    reminders = db.session.scalars(
        db.select(Reminder)
        .where(
            Reminder.church_id == current_user.church_id
        )
        .order_by(
            Reminder.scheduled_at.desc()
        )
    ).all()

    return render_template(
        "reminders/index.html",
        reminders=reminders
    )


# =========================================================
# REMINDER DETAILS
# =========================================================

@dashboard_bp.route(
    "/reminders/<int:reminder_id>"
)
@login_required
def reminder_details(reminder_id):

    reminder = db.session.scalar(
        db.select(Reminder)
        .where(
            Reminder.id == reminder_id,
            Reminder.church_id == current_user.church_id
        )
    )

    if reminder is None:

        flash(
            "Reminder not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.reminders")
        )

    return render_template(
        "reminders/details.html",
        reminder=reminder
    )


# =========================================================
# EDIT REMINDER
# =========================================================

@dashboard_bp.route(
    "/reminders/<int:reminder_id>/edit",
    methods=["GET", "POST"]
)
@login_required
def edit_reminder(reminder_id):

    reminder = db.session.scalar(
        db.select(Reminder)
        .where(
            Reminder.id == reminder_id,
            Reminder.church_id == current_user.church_id
        )
    )

    if reminder is None:

        flash(
            "Reminder not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.reminders")
        )

    programs, groups = get_reminder_form_data()

    if request.method == "POST":

        title = request.form.get(
            "title",
            ""
        ).strip()

        message = request.form.get(
            "message",
            ""
        ).strip()

        program_id = request.form.get(
            "program_id",
            ""
        ).strip()

        audience_type = request.form.get(
            "audience_type",
            "all"
        ).strip()

        scheduled_at_raw = request.form.get(
            "scheduled_at",
            ""
        ).strip()

        recurrence_type = request.form.get(
            "recurrence_type",
            "one_time"
        ).strip()

        recurrence_interval_raw = request.form.get(
            "recurrence_interval",
            "1"
        ).strip()

        recurrence_end_date_raw = request.form.get(
            "recurrence_end_date",
            ""
        ).strip()

        group_ids = request.form.getlist(
            "group_ids"
        )

        if not title:

            flash(
                "Please enter a reminder title.",
                "error"
            )

            return render_template(
                "reminders/edit.html",
                reminder=reminder,
                programs=programs,
                groups=groups
            )

        if not message:

            flash(
                "Please enter the SMS message.",
                "error"
            )

            return render_template(
                "reminders/edit.html",
                reminder=reminder,
                programs=programs,
                groups=groups
            )

        if not scheduled_at_raw:

            flash(
                "Please select a date and time.",
                "error"
            )

            return render_template(
                "reminders/edit.html",
                reminder=reminder,
                programs=programs,
                groups=groups
            )

        try:

            scheduled_at = datetime.strptime(
                scheduled_at_raw,
                "%Y-%m-%dT%H:%M"
            )

        except ValueError:

            flash(
                "Invalid date and time format.",
                "error"
            )

            return render_template(
                "reminders/edit.html",
                reminder=reminder,
                programs=programs,
                groups=groups
            )

        allowed_recurrence_types = {
            "one_time",
            "daily",
            "weekly",
            "monthly",
            "yearly",
            "interval_days",
            "interval_weeks",
            "interval_months",
            "interval_years"
        }

        if recurrence_type not in allowed_recurrence_types:

            flash(
                "Invalid recurrence option selected.",
                "error"
            )

            return render_template(
                "reminders/edit.html",
                reminder=reminder,
                programs=programs,
                groups=groups
            )

        try:

            recurrence_interval = int(
                recurrence_interval_raw or 1
            )

        except ValueError:

            flash(
                "Recurrence interval must be a number.",
                "error"
            )

            return render_template(
                "reminders/edit.html",
                reminder=reminder,
                programs=programs,
                groups=groups
            )

        if recurrence_interval < 1:

            flash(
                "Recurrence interval must be at least 1.",
                "error"
            )

            return render_template(
                "reminders/edit.html",
                reminder=reminder,
                programs=programs,
                groups=groups
            )

        recurrence_end_date = None

        if recurrence_end_date_raw:

            try:

                recurrence_end_date = datetime.strptime(
                    recurrence_end_date_raw,
                    "%Y-%m-%d"
                )

            except ValueError:

                flash(
                    "Invalid recurrence end date.",
                    "error"
                )

                return render_template(
                    "reminders/edit.html",
                    reminder=reminder,
                    programs=programs,
                    groups=groups
                )

            if recurrence_end_date < scheduled_at:

                flash(
                    "The recurrence end date cannot be before the scheduled date.",
                    "error"
                )

                return render_template(
                    "reminders/edit.html",
                    reminder=reminder,
                    programs=programs,
                    groups=groups
                )

        if recurrence_type == "one_time":

            recurrence_interval = 1
            recurrence_end_date = None

        allowed_audiences = {
            "all",
            "groups"
        }

        if audience_type not in allowed_audiences:

            flash(
                "Invalid audience selected.",
                "error"
            )

            return render_template(
                "reminders/edit.html",
                reminder=reminder,
                programs=programs,
                groups=groups
            )

        selected_groups = []

        if audience_type == "groups":

            if not group_ids:

                flash(
                    "Please select at least one group.",
                    "error"
                )

                return render_template(
                    "reminders/edit.html",
                    reminder=reminder,
                    programs=programs,
                    groups=groups
                )

            selected_group_ids = []

            for group_id in group_ids:

                try:

                    selected_group_ids.append(
                        int(group_id)
                    )

                except ValueError:

                    continue

            if not selected_group_ids:

                flash(
                    "Invalid group selection.",
                    "error"
                )

                return render_template(
                    "reminders/edit.html",
                    reminder=reminder,
                    programs=programs,
                    groups=groups
                )

            selected_groups = db.session.scalars(
                db.select(Group)
                .where(
                    Group.id.in_(
                        selected_group_ids
                    ),
                    Group.church_id == current_user.church_id,
                    Group.is_active.is_(True)
                )
            ).all()

            if len(selected_groups) != len(
                set(selected_group_ids)
            ):

                flash(
                    "One or more selected groups are invalid.",
                    "error"
                )

                return render_template(
                    "reminders/edit.html",
                    reminder=reminder,
                    programs=programs,
                    groups=groups
                )

        program = None

        if program_id:

            try:

                program_id_int = int(
                    program_id
                )

            except ValueError:

                flash(
                    "Invalid program selected.",
                    "error"
                )

                return render_template(
                    "reminders/edit.html",
                    reminder=reminder,
                    programs=programs,
                    groups=groups
                )

            program = db.session.scalar(
                db.select(Program)
                .where(
                    Program.id == program_id_int,
                    Program.church_id == current_user.church_id,
                    Program.is_active.is_(True)
                )
            )

            if not program:

                flash(
                    "The selected program is invalid.",
                    "error"
                )

                return render_template(
                    "reminders/edit.html",
                    reminder=reminder,
                    programs=programs,
                    groups=groups
                )

        # -------------------------------------------------
        # Update reminder
        # -------------------------------------------------

        reminder.title = title
        reminder.message = message
        reminder.program_id = (
            program.id
            if program
            else None
        )
        reminder.audience_type = audience_type
        reminder.scheduled_at = scheduled_at
        reminder.recurrence_type = recurrence_type
        reminder.recurrence_interval = recurrence_interval
        reminder.recurrence_end_date = recurrence_end_date

        if audience_type == "groups":

            reminder.groups = selected_groups

        else:

            reminder.groups = []

        # -------------------------------------------------
        # Recalculate next run
        # -------------------------------------------------

        if recurrence_type == "one_time":

            reminder.next_run_at = None

        else:

            reminder.next_run_at = calculate_next_future_run(
                scheduled_at=scheduled_at,
                recurrence_type=recurrence_type,
                recurrence_interval=recurrence_interval
            )

            if (
                reminder.next_run_at
                and recurrence_end_date
                and reminder.next_run_at > recurrence_end_date
            ):

                reminder.next_run_at = None

        # -------------------------------------------------
        # Editing means the reminder is ready to be
        # scheduled again.
        # -------------------------------------------------

        reminder.status = "scheduled"
        reminder.is_active = True
        reminder.failure_reason = None

        db.session.commit()

        flash(
            "Reminder updated successfully.",
            "success"
        )

        return redirect(
            url_for(
                "dashboard.reminder_details",
                reminder_id=reminder.id
            )
        )

    return render_template(
        "reminders/edit.html",
        reminder=reminder,
        programs=programs,
        groups=groups
    )


# =========================================================
# TOGGLE REMINDER STATUS
# =========================================================

@dashboard_bp.route(
    "/reminders/<int:reminder_id>/toggle-status",
    methods=["POST"]
)
@login_required
def toggle_reminder_status(reminder_id):

    reminder = db.session.scalar(
        db.select(Reminder)
        .where(
            Reminder.id == reminder_id,
            Reminder.church_id == current_user.church_id
        )
    )

    if reminder is None:

        flash(
            "Reminder not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.reminders")
        )

    if reminder.is_active:

        reminder.is_active = False

        reminder.status = "paused"

        flash(
            f'"{reminder.title}" has been paused.',
            "success"
        )

    else:

        reminder.is_active = True

        reminder.status = "scheduled"

        if reminder.recurrence_type == "one_time":

            reminder.next_run_at = None

        else:

            reminder.next_run_at = calculate_next_future_run(
                scheduled_at=reminder.scheduled_at,
                recurrence_type=reminder.recurrence_type,
                recurrence_interval=reminder.recurrence_interval
            )

            if (
                reminder.next_run_at
                and reminder.recurrence_end_date
                and reminder.next_run_at > reminder.recurrence_end_date
            ):

                reminder.next_run_at = None

        reminder.failure_reason = None

        flash(
            f'"{reminder.title}" has been activated.',
            "success"
        )

    db.session.commit()

    return redirect(
        url_for(
            "dashboard.reminder_details",
            reminder_id=reminder.id
        )
    )


# =========================================================
# DELETE REMINDER
# =========================================================

@dashboard_bp.route(
    "/reminders/<int:reminder_id>/delete",
    methods=["POST"]
)
@login_required
def delete_reminder(reminder_id):

    reminder = db.session.scalar(
        db.select(Reminder)
        .where(
            Reminder.id == reminder_id,
            Reminder.church_id == current_user.church_id
        )
    )

    if reminder is None:

        flash(
            "Reminder not found.",
            "error"
        )

        return redirect(
            url_for("dashboard.reminders")
        )

    reminder_title = reminder.title

    db.session.delete(
        reminder
    )

    db.session.commit()

    flash(
        f'"{reminder_title}" was deleted successfully.',
        "success"
    )

    return redirect(
        url_for(
            "dashboard.reminders"
        )
    )

# =========================================================
# MESSAGES
# =========================================================

@dashboard_bp.route("/messages")
@login_required
def messages():

    search = request.args.get(
        "search",
        ""
    ).strip()

    status = request.args.get(
        "status",
        ""
    ).strip()

    source = request.args.get(
        "source",
        ""
    ).strip()

    query = db.select(Message).where(
        Message.church_id == current_user.church_id
    )

    # -------------------------------------------------
    # SEARCH
    # -------------------------------------------------

    if search:

        search_pattern = f"%{search}%"

        query = query.where(
            db.or_(
                Message.recipient_phone.ilike(
                    search_pattern
                ),
                Message.message.ilike(
                    search_pattern
                )
            )
        )

    # -------------------------------------------------
    # STATUS FILTER
    # -------------------------------------------------

    allowed_statuses = {
        "pending",
        "sent",
        "failed"
    }

    if status in allowed_statuses:

        query = query.where(
            Message.status == status
        )

    # -------------------------------------------------
    # SOURCE FILTER
    # -------------------------------------------------

    allowed_sources = {
        "campaign",
        "reminder",
        "manual"
    }

    if source in allowed_sources:

        query = query.where(
            Message.source == source
        )

    # -------------------------------------------------
    # ORDER
    # -------------------------------------------------

    query = query.order_by(
        Message.created_at.desc()
    )

    messages_list = db.session.scalars(
        query
    ).all()

    # -------------------------------------------------
    # STATISTICS
    # -------------------------------------------------

    total_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id
        )
    ) or 0

    sent_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id,
            Message.status == "sent"
        )
    ) or 0

    failed_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id,
            Message.status == "failed"
        )
    ) or 0

    pending_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id,
            Message.status == "pending"
        )
    ) or 0

    return render_template(
        "messages/index.html",
        messages=messages_list,
        search=search,
        selected_status=status,
        selected_source=source,
        total_messages=total_messages,
        sent_messages=sent_messages,
        failed_messages=failed_messages,
        pending_messages=pending_messages
    )

    # =========================================================
# REPORTS
# =========================================================

@dashboard_bp.route("/reports")
@login_required
def reports():

    # -----------------------------------------------------
    # DATE FILTER
    # -----------------------------------------------------

    start_date = request.args.get(
        "start_date",
        ""
    ).strip()

    end_date = request.args.get(
        "end_date",
        ""
    ).strip()

    # -----------------------------------------------------
    # BASE MESSAGE QUERY
    # -----------------------------------------------------

    message_query = db.select(Message).where(
        Message.church_id == current_user.church_id
    )

    # Apply start date
    if start_date:

        try:

            start_datetime = datetime.strptime(
                start_date,
                "%Y-%m-%d"
            )

            message_query = message_query.where(
                Message.created_at >= start_datetime
            )

        except ValueError:

            start_date = ""

    # Apply end date
    if end_date:

        try:

            end_datetime = datetime.strptime(
                end_date,
                "%Y-%m-%d"
            )

            # Include the entire selected end date
            end_datetime = (
                end_datetime + timedelta(days=1)
            )

            message_query = message_query.where(
                Message.created_at < end_datetime
            )

        except ValueError:

            end_date = ""

    # -----------------------------------------------------
    # MESSAGE STATISTICS
    # -----------------------------------------------------

    total_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id,
            *(
                [
                    Message.created_at >= start_datetime
                ]
                if start_date
                else []
            ),
            *(
                [
                    Message.created_at < end_datetime
                ]
                if end_date
                else []
            )
        )
    ) or 0

    sent_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id,
            Message.status == "sent",
            *(
                [
                    Message.created_at >= start_datetime
                ]
                if start_date
                else []
            ),
            *(
                [
                    Message.created_at < end_datetime
                ]
                if end_date
                else []
            )
        )
    ) or 0

    failed_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id,
            Message.status == "failed",
            *(
                [
                    Message.created_at >= start_datetime
                ]
                if start_date
                else []
            ),
            *(
                [
                    Message.created_at < end_datetime
                ]
                if end_date
                else []
            )
        )
    ) or 0

    pending_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id,
            Message.status == "pending",
            *(
                [
                    Message.created_at >= start_datetime
                ]
                if start_date
                else []
            ),
            *(
                [
                    Message.created_at < end_datetime
                ]
                if end_date
                else []
            )
        )
    ) or 0

    # -----------------------------------------------------
    # SOURCE STATISTICS
    # -----------------------------------------------------

    campaign_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id,
            Message.source == "campaign",
            *(
                [
                    Message.created_at >= start_datetime
                ]
                if start_date
                else []
            ),
            *(
                [
                    Message.created_at < end_datetime
                ]
                if end_date
                else []
            )
        )
    ) or 0

    reminder_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id,
            Message.source == "reminder",
            *(
                [
                    Message.created_at >= start_datetime
                ]
                if start_date
                else []
            ),
            *(
                [
                    Message.created_at < end_datetime
                ]
                if end_date
                else []
            )
        )
    ) or 0

    manual_messages = db.session.scalar(
        db.select(
            db.func.count(Message.id)
        ).where(
            Message.church_id == current_user.church_id,
            Message.source == "manual",
            *(
                [
                    Message.created_at >= start_datetime
                ]
                if start_date
                else []
            ),
            *(
                [
                    Message.created_at < end_datetime
                ]
                if end_date
                else []
            )
        )
    ) or 0

    # -----------------------------------------------------
    # MEMBER STATISTICS
    # -----------------------------------------------------

    total_members = db.session.scalar(
        db.select(
            db.func.count(Member.id)
        ).where(
            Member.church_id == current_user.church_id
        )
    ) or 0

    active_members = db.session.scalar(
        db.select(
            db.func.count(Member.id)
        ).where(
            Member.church_id == current_user.church_id,
            Member.is_active.is_(True)
        )
    ) or 0

    sms_opted_members = db.session.scalar(
        db.select(
            db.func.count(Member.id)
        ).where(
            Member.church_id == current_user.church_id,
            Member.is_active.is_(True),
            Member.sms_opt_in.is_(True)
        )
    ) or 0

    # -----------------------------------------------------
    # CAMPAIGN STATISTICS
    # -----------------------------------------------------

    total_campaigns = db.session.scalar(
        db.select(
            db.func.count(Campaign.id)
        ).where(
            Campaign.church_id == current_user.church_id
        )
    ) or 0

    sent_campaigns = db.session.scalar(
        db.select(
            db.func.count(Campaign.id)
        ).where(
            Campaign.church_id == current_user.church_id,
            Campaign.status == "sent"
        )
    ) or 0

    scheduled_campaigns = db.session.scalar(
        db.select(
            db.func.count(Campaign.id)
        ).where(
            Campaign.church_id == current_user.church_id,
            Campaign.status == "scheduled"
        )
    ) or 0

    # -----------------------------------------------------
    # REMINDER STATISTICS
    # -----------------------------------------------------

    total_reminders = db.session.scalar(
        db.select(
            db.func.count(Reminder.id)
        ).where(
            Reminder.church_id == current_user.church_id
        )
    ) or 0

    active_reminders = db.session.scalar(
        db.select(
            db.func.count(Reminder.id)
        ).where(
            Reminder.church_id == current_user.church_id,
            Reminder.is_active.is_(True)
        )
    ) or 0

    completed_reminders = db.session.scalar(
        db.select(
            db.func.count(Reminder.id)
        ).where(
            Reminder.church_id == current_user.church_id,
            Reminder.status == "completed"
        )
    ) or 0

    # -----------------------------------------------------
    # DELIVERY RATE
    # -----------------------------------------------------

    delivery_rate = 0

    if total_messages > 0:

        delivery_rate = round(
            (sent_messages / total_messages) * 100,
            1
        )

    # -----------------------------------------------------
    # RECENT MESSAGES
    # -----------------------------------------------------

    recent_messages_query = db.select(Message).where(
        Message.church_id == current_user.church_id
    )

    if start_date:

        recent_messages_query = recent_messages_query.where(
            Message.created_at >= start_datetime
        )

    if end_date:

        recent_messages_query = recent_messages_query.where(
            Message.created_at < end_datetime
        )

    recent_messages = db.session.scalars(
        recent_messages_query
        .order_by(
            Message.created_at.desc()
        )
        .limit(10)
    ).all()

    # -----------------------------------------------------
    # RENDER
    # -----------------------------------------------------

    return render_template(
        "reports/index.html",

        start_date=start_date,
        end_date=end_date,

        total_messages=total_messages,
        sent_messages=sent_messages,
        failed_messages=failed_messages,
        pending_messages=pending_messages,

        campaign_messages=campaign_messages,
        reminder_messages=reminder_messages,
        manual_messages=manual_messages,

        total_members=total_members,
        active_members=active_members,
        sms_opted_members=sms_opted_members,

        total_campaigns=total_campaigns,
        sent_campaigns=sent_campaigns,
        scheduled_campaigns=scheduled_campaigns,

        total_reminders=total_reminders,
        active_reminders=active_reminders,
        completed_reminders=completed_reminders,

        delivery_rate=delivery_rate,

        recent_messages=recent_messages
    )

# =========================================================
# SETTINGS
# =========================================================

@dashboard_bp.route("/settings", methods=["GET", "POST"])
@login_required
def settings():

    church = db.session.get(
        Church,
        current_user.church_id
    )

    if not church:

        flash(
            "Church information could not be found.",
            "error"
        )

        return redirect(
            url_for("dashboard.home")
        )

    if request.method == "POST":

        church.name = request.form.get(
            "name",
            ""
        ).strip()

        church.location = request.form.get(
            "location",
            ""
        ).strip()

        church.phone = request.form.get(
            "phone",
            ""
        ).strip()

        church.email = request.form.get(
            "email",
            ""
        ).strip()

        if not church.name:

            flash(
                "Church name is required.",
                "error"
            )

            return render_template(
                "settings/index.html",
                church=church
            )

        db.session.commit()

        flash(
            "Church settings updated successfully.",
            "success"
        )

        return redirect(
            url_for("dashboard.settings")
        )

    return render_template(
        "settings/index.html",
        church=church
    )